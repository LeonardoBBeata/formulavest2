from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Formulavest_Documentacao_TCC.docx"

INK = RGBColor(15, 23, 42)
MUTED = RGBColor(100, 116, 139)
LIGHT_FILL = "F3F4F6"
MID_FILL = "E5E7EB"
BORDER = "D7DBE2"
HEADER_FILL = "1F2937"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_grid = tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.append(tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        tbl_grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(rows, widths, header=True):
    doc = table.doc
    tbl = doc.add_table(rows=len(rows), cols=len(widths))
    set_table_width(tbl, widths)
    set_table_borders(tbl)
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = tbl.rows[r_idx].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(text))
            set_run_font(
                run,
                size=9.5 if r_idx else 10,
                color=RGBColor(255, 255, 255) if header and r_idx == 0 else INK,
                bold=(header and r_idx == 0),
            )
            if header and r_idx == 0:
                set_cell_shading(cell, HEADER_FILL)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, LIGHT_FILL)
    doc.add_paragraph()
    return tbl


def add_para(doc, text="", style=None, bold=False, italic=False, color=INK, size=11, align=None, after=6):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        run = p.add_run(item)
        set_run_font(run, size=11, color=INK)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        run = p.add_run(item)
        set_run_font(run, size=11, color=INK)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12 if level == 2 else 8)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6 if level == 2 else 4)
    for run in p.runs:
        set_run_font(
            run,
            size=16 if level == 1 else 13 if level == 2 else 12,
            color=INK,
            bold=True,
        )
    return p


def page_break(doc):
    doc.add_page_break()


def configure_document(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    title_style = styles["Title"]
    title_style.font.name = "Calibri"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title_style.font.size = Pt(22)
    title_style.font.color.rgb = INK
    title_style.font.bold = True

    for name, size, color in (
        ("Heading 1", 16, INK),
        ("Heading 2", 13, INK),
        ("Heading 3", 12, INK),
    ):
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True

    header = sec.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("FormulaVest | Documento técnico do projeto")
    set_run_font(run, size=9, color=MUTED)

    footer = sec.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("FormulaVest | Documentação técnica e operacional")
    set_run_font(run, size=9, color=MUTED)

    props = doc.core_properties
    props.title = "FormulaVest — Documentação Técnica e Operacional"
    props.subject = "Arquitetura, operação e requisitos da plataforma FormulaVest"
    props.author = "Equipe FormulaVest"
    props.keywords = "educação, simulados, provas online, Node.js, PostgreSQL"


def cover(doc):
    add_para(doc, "FORMULAVEST", bold=True, size=14, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    add_para(
        doc,
        "DOCUMENTAÇÃO TÉCNICA E OPERACIONAL",
        style="Title",
        bold=True,
        size=22,
        color=INK,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=8,
    )
    add_para(
        doc,
        "Plataforma Web de Estudos, Simulados, Provas ao Vivo e Gestao Escolar",
        size=14,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=28,
    )
    rows = [
        ("Tipo de documento", "Relatorio tecnico em formato de TCC"),
        ("Sistema", "FormulaVest"),
        ("Stack principal", "Node.js, Express, PostgreSQL, HTML, CSS e JavaScript"),
        ("Versao documentada", date.today().strftime("%d/%m/%Y")),
        ("Escopo", "Arquitetura, requisitos, banco de dados, funcionalidades, testes, implantacao e manutencao"),
    ]
    table(rows, [2600, 6760], header=False)
    add_para(
        doc,
        "Documento elaborado para registrar a estrutura tecnica, a proposta pedagogica e os fluxos operacionais do sistema FormulaVest.",
        italic=True,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=18,
    )
    add_para(doc, "Brasil, 2026", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    page_break(doc)


def add_note(doc, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.1
    lead = p.add_run(f"{title}. ")
    set_run_font(lead, size=11, color=INK, bold=True)
    text = p.add_run(body)
    set_run_font(text, size=11, color=INK)


def build_doc():
    doc = Document()
    configure_document(doc)
    table.doc = doc
    cover(doc)

    heading(doc, "Resumo", 1)
    add_para(
        doc,
        "O FormulaVest e uma plataforma web voltada ao apoio de estudos, simulados e gestao de atividades avaliativas. "
        "O sistema integra area do aluno, painel do professor, painel administrativo, provas geradas por inteligencia artificial, "
        "provas ao vivo com codigo de acesso, ranking, conquistas, redacao, relatorios e administracao de empresas, escolas, periodos e salas.",
    )
    add_para(
        doc,
        "A documentação descreve objetivos, requisitos, arquitetura, dados, fluxos de usuário, segurança, testes, operação e manutenção. "
        "Também estabelece critérios de qualidade e rastreabilidade para que a solução possa ser apresentada, implantada e evoluída de forma segura.",
    )
    heading(doc, "Palavras-chave", 2)
    add_para(doc, "Educacao; simulados; Node.js; PostgreSQL; provas online; gamificacao; FormulaVest.")
    page_break(doc)

    heading(doc, "Sumario", 1)
    for item in [
        "1. Introducao",
        "2. Objetivos",
        "3. Justificativa",
        "4. Metodologia e tecnologias",
        "5. Arquitetura do sistema",
        "6. Modelo de dados",
        "7. Funcionalidades por perfil",
        "8. Fluxos principais",
        "9. Segurança, privacidade e integridade",
        "10. Requisitos e regras de negócio",
        "11. Testes e validação",
        "12. Implantação e operação",
        "13. Manutenção evolutiva",
        "14. Conclusão",
        "Apendices",
    ]:
        add_para(doc, item, after=3)
    page_break(doc)

    heading(doc, "1. Introducao", 1)
    add_para(
        doc,
        "O FormulaVest foi desenvolvido como uma solucao web para apoiar estudantes e professores em rotinas de preparacao para provas. "
        "A aplicacao centraliza simulados, historico de desempenho, ranking, conquistas, redacao e provas criadas por professores. "
        "A proposta combina estudo individual com interacoes de turma, permitindo que professores transmitam provas ao vivo por codigo e acompanhem resultados em tempo real.",
    )
    add_para(
        doc,
        "A necessidade do projeto surge da dificuldade de organizar estudos, aplicar atividades digitais e medir desempenho com rapidez. "
        "Ao reunir geracao de questoes, dashboards e recursos administrativos em uma unica plataforma, o sistema reduz a fragmentacao de ferramentas e facilita o acompanhamento pedagogico.",
    )

    heading(doc, "2. Objetivos", 1)
    heading(doc, "2.1 Objetivo geral", 2)
    add_para(
        doc,
        "Construir uma plataforma educacional web capaz de gerar, aplicar, corrigir e acompanhar simulados e provas, oferecendo experiencias especificas para alunos, professores e administradores.",
    )
    heading(doc, "2.2 Objetivos especificos", 2)
    add_bullets(doc, [
        "Permitir que alunos realizem simulados ENEM e Provao Paulista com registro de desempenho.",
        "Oferecer ao professor um ambiente para criar provas manuais, gerar questoes por IA, importar arquivos e transmitir provas ao vivo.",
        "Disponibilizar codigo de acesso para provas de turma e placar em tempo real.",
        "Organizar usuarios por empresa, escola, periodo e sala.",
        "Gerar relatorios em PDF e XLSX para acompanhamento pedagogico.",
        "Aplicar recursos de gamificacao, como XP, nivel, ranking e conquistas.",
        "Manter autenticacao, autorizacao por perfil e controle de usuarios banidos.",
    ])

    heading(doc, "3. Justificativa", 1)
    add_para(
        doc,
        "Ferramentas educacionais eficientes precisam combinar usabilidade, rastreabilidade e rapidez. Em ambientes escolares, a aplicacao de provas digitais deve ser simples para o professor e objetiva para o aluno. "
        "O FormulaVest atende a esse cenario ao fornecer telas especializadas, dashboards e uma camada administrativa capaz de refletir a estrutura real da instituicao.",
    )
    add_note(
        doc,
        "Contribuicao do projeto",
        "A principal contribuição é integrar estudo individual, prova ao vivo e gestão escolar em uma mesma base de dados, reduzindo retrabalho e viabilizando análises posteriores de desempenho.",
    )

    heading(doc, "4. Metodologia e tecnologias", 1)
    add_para(
        doc,
        "O projeto utiliza uma arquitetura web tradicional, com servidor Node.js/Express, persistencia em PostgreSQL e interface publica composta por HTML, CSS e JavaScript puro. "
        "A abordagem favorece simplicidade de deploy, baixo acoplamento entre telas e facilidade de manutencao por arquivos estaticos.",
    )
    table([
        ("Camada", "Tecnologia", "Responsabilidade"),
        ("Frontend", "HTML, CSS, JavaScript", "Telas do aluno, professor, admin, login, perfil e PWA"),
        ("Backend", "Node.js + Express", "Rotas REST, autenticacao, regras de negocio e arquivos estaticos"),
        ("Banco", "PostgreSQL", "Usuarios, provas, empresas, escolas, salas, respostas e conquistas"),
        ("IA", "OpenRouter API", "Geracao e correcao assistida por IA"),
        ("Relatorios", "PDFKit e ExcelJS", "Exportacao de PDFs e planilhas XLSX"),
        ("Infra", "Docker, docker-compose, GitHub Actions", "Ambiente local, banco e CI basico"),
    ], [1800, 2300, 5260])

    heading(doc, "5. Arquitetura do sistema", 1)
    add_para(
        doc,
        "A aplicacao e organizada por responsabilidades. O arquivo server.js configura middlewares globais, seguranca, CORS, JSON, arquivos estaticos e registra modulos de rotas. "
        "As rotas ficam em routes/, servicos externos em services/, configuracoes em config/ e interfaces estaticas em public/.",
    )
    table([
        ("Modulo", "Descricao"),
        ("server.js", "Inicializacao do Express, middlewares, arquivos estaticos e registro das rotas."),
        ("config/database.js", "Conexao PostgreSQL, criacao do schema inicial e usuario master."),
        ("middlewares/auth.js", "JWT, payload do usuario e middleware de permissao por perfil."),
        ("routes/auth.js", "Registro, verificacao de email, login, refresh token e reset de senha."),
        ("routes/provas.js", "Simulados, provas de professor, prova ao vivo, respostas, exports e PDF."),
        ("routes/admin.js", "Gestao de usuarios, alunos, professores, estrutura escolar, empresas e estatisticas."),
        ("routes/user.js", "Perfil, XP, conquistas, upload de foto e dados do usuario logado."),
        ("public/", "Telas e scripts do aluno, professor, admin, login, perfil e PWA."),
    ], [2600, 6760])
    add_note(
        doc,
        "Fluxo macro",
        "Navegador -> arquivos public/ -> fetch para rotas Express -> validacao JWT -> regras de permissao -> consultas PostgreSQL -> resposta JSON ou arquivo exportado.",
    )

    heading(doc, "6. Modelo de dados", 1)
    add_para(
        doc,
        "O modelo de dados representa usuarios, estrutura institucional, provas individuais, provas de professor e respostas. "
        "A separacao entre provas_ativas, provas e provas_professor permite distinguir simulados individuais de avaliacoes criadas por professores.",
    )
    table([
        ("Tabela", "Finalidade principal"),
        ("usuarios", "Conta, email, senha, verificacao, XP, nivel, papel, empresa, escola, periodo e sala."),
        ("empresas", "Organizacao mantenedora de escolas."),
        ("escolas", "Unidades vinculadas a uma empresa."),
        ("periodos", "Agrupamentos escolares dentro da escola."),
        ("salas", "Turmas vinculadas a periodos."),
        ("provas_ativas", "Simulados individuais ainda em andamento."),
        ("provas", "Historico de simulados finalizados pelo aluno."),
        ("provas_professor", "Provas criadas por professores, com codigo, status e questoes JSONB."),
        ("respostas_provas_professor", "Respostas de alunos em provas de professor, com nota e marcador finalizada."),
        ("provas_professor_participantes", "Controle de entrada e limite de alunos por prova ao vivo."),
        ("conquistas_historico", "Registro de conquistas obtidas pelos alunos."),
        ("refresh_tokens", "Sessao persistente por cookie httpOnly."),
    ], [2600, 6760])

    heading(doc, "6.1 Regras importantes de dados", 2)
    add_bullets(doc, [
        "Usuarios possuem role para controle de autorizacao: formulavest_master, empresa_admin, diretor, coordenador, professor e aluno.",
        "Provas de professor armazenam questoes em JSONB para permitir importacao e geracao flexivel.",
        "Respostas parciais da prova ao vivo sao mantidas sem bloquear a finalizacao definitiva.",
        "O limite de participantes e controlado por max_alunos e pela tabela provas_professor_participantes.",
        "O historico de conquistas evita duplicidade e permite exibir marcos de evolucao.",
    ])

    heading(doc, "7. Funcionalidades por perfil", 1)
    table([
        ("Perfil", "Funcionalidades"),
        ("Aluno", "Gerar simulados, responder provas com codigo, visualizar ranking, XP, conquistas, redacao e perfil."),
        ("Professor", "Criar/importar provas, gerar questoes por IA, transmitir prova ao vivo, acompanhar placar e exportar relatorios."),
        ("Coordenador/Diretor", "Gerenciar alunos/professores dentro da escola, estrutura e acompanhamento."),
        ("Admin da empresa", "Gerenciar escolas, usuarios e dados da empresa."),
        ("Master", "Gerenciar empresas, admins e visao global da plataforma."),
    ], [1900, 7460])

    heading(doc, "7.1 Experiência do aluno e indicadores", 2)
    add_para(
        doc,
        "Após concluir uma prova, o aluno recebe uma tela de resultado com nota, total de acertos, percentual, tempo empregado, XP recebido e evolução de nível. "
        "O painel de estatísticas consolida histórico de tentativas, gráfico de desempenho e comparação por matéria. Esses elementos transformam a correção em uma devolutiva pedagógica clara, e não apenas em uma nota isolada.",
    )
    table([
        ("Indicador", "Uso pedagógico"),
        ("Acertos e percentual", "Permite identificar domínio global da prova."),
        ("Desempenho por matéria", "Orienta revisão focada em áreas com menor aproveitamento."),
        ("XP e nível", "Dá retorno imediato de progresso e incentiva recorrência."),
        ("Ranking", "Cria comparação opcional e motivação em contexto de turma."),
        ("Conquistas", "Reconhece marcos como primeira prova, sequência e evolução de nível."),
    ], [2800, 6560])
    heading(doc, "7.2 Conquistas implementadas", 2)
    add_bullets(doc, [
        "Primeira prova, três provas, dez provas e cinquenta provas concluídas.",
        "Prova perfeita, para aproveitamento total em uma tentativa finalizada.",
        "Marcos de XP e de nível, incluindo os patamares cinco e dez.",
        "Sequências de estudo de três e sete dias, calculadas a partir de tentativas finalizadas.",
        "Notificação não intrusiva no canto superior direito quando uma nova conquista é registrada.",
    ])

    heading(doc, "8. Fluxos principais", 1)
    heading(doc, "8.1 Geracao de simulado individual", 2)
    add_numbered(doc, [
        "Aluno acessa a tela principal e escolhe ENEM ou Provao Paulista.",
        "O frontend chama /gerar-enem ou /gerar-provao com JWT.",
        "O backend solicita questoes a IA e salva a prova em provas_ativas.",
        "O frontend renderiza as questoes e habilita o botao de finalizar.",
        "Ao finalizar, /salvar-prova corrige pelo gabarito salvo, registra provas e atualiza XP/nivel.",
    ])
    heading(doc, "8.2 Prova ao vivo", 2)
    add_numbered(doc, [
        "Professor cria uma prova e escolhe a sala e limite de alunos.",
        "Professor clica em Transmitir; a prova recebe codigo, status ativa e horarios de inicio/fim.",
        "Aluno entra pelo codigo; o sistema valida status, sala, limite de participantes e tempo.",
        "Aluno responde pergunta a pergunta; as respostas parciais atualizam o placar.",
        "No final, a resposta e marcada como finalizada, XP e atualizado e o professor pode exportar resultados.",
    ])
    heading(doc, "8.3 Administracao", 2)
    add_numbered(doc, [
        "Admin entra no painel administrativo.",
        "A tela carrega perfil, estatisticas, usuarios, provas e estrutura escolar.",
        "Operacoes de criacao/edicao usam rotas protegidas por role.",
        "Professores podem ser associados a salas para limitar seu escopo operacional.",
    ])

    heading(doc, "9. Seguranca, privacidade e integridade", 1)
    add_bullets(doc, [
        "Autenticacao baseada em JWT no header Authorization.",
        "Refresh token armazenado em cookie httpOnly para sessao persistente.",
        "Senhas com hash bcrypt.",
        "Middleware permitir(...) controla acesso por perfil.",
        "CORS restrito a origens configuradas.",
        "Helmet, rate limit e limites de payload reduzem riscos comuns.",
        "Uploads de prova aceitam apenas PDF e DOCX com limite de tamanho.",
        "Alunos banidos sao bloqueados no login.",
    ])
    add_note(
        doc,
        "Ponto de atencao",
        "A segurança depende de JWT_SECRET forte, MASTER_PASSWORD seguro, uso de HTTPS em produção e controle adequado das variáveis .env.",
    )

    heading(doc, "10. Requisitos e regras de negócio", 1)
    heading(doc, "10.1 Requisitos funcionais", 2)
    table([
        ("ID", "Requisito", "Critério de aceite"),
        ("RF01", "Cadastrar e autenticar usuários", "Conta é criada, verificada e recebe sessão válida."),
        ("RF02", "Gerar simulados", "Questões são apresentadas e a tentativa fica associada ao aluno."),
        ("RF03", "Finalizar e corrigir prova", "Resultado, acertos, XP, nível e histórico são persistidos."),
        ("RF04", "Aplicar prova ao vivo", "Código, horário, sala e limite de participantes são validados."),
        ("RF05", "Administrar estrutura", "Perfis autorizados gerenciam empresa, escola, período e sala."),
        ("RF06", "Produzir relatórios", "Professor exporta dados de suas provas e turmas em PDF ou XLSX."),
        ("RF07", "Gamificar progresso", "Conquistas únicas são registradas e notificadas no cliente."),
    ], [900, 3600, 4860])
    heading(doc, "10.2 Requisitos não funcionais", 2)
    table([
        ("Categoria", "Diretriz adotada"),
        ("Segurança", "Autorização por papel, JWT, bcrypt, rate limit, CORS e validação de entrada."),
        ("Usabilidade", "Telas distintas por perfil, feedback de carregamento e estados de erro compreensíveis."),
        ("Confiabilidade", "Registro transacional de resultados, validação de prova ativa e tratamento centralizado de erros."),
        ("Desempenho", "Compressão HTTP, cache control para estáticos, cache de aplicação e consultas segmentadas."),
        ("Compatibilidade", "Aplicação web responsiva e PWA com service worker para ativos seguros."),
        ("Manutenibilidade", "Rotas organizadas por domínio, configuração por ambiente e scripts de banco versionáveis."),
    ], [2500, 6860])
    heading(doc, "10.3 Regras de negócio críticas", 2)
    add_bullets(doc, [
        "Uma prova de professor só pode receber respostas enquanto estiver ativa e dentro de sua janela de aplicação.",
        "O aluno só pode ingressar se pertencer à sala permitida e houver capacidade disponível.",
        "A pontuação de uma tentativa finalizada não deve ser recalculada a partir de respostas parciais.",
        "A conquista é identificada por chave e só pode ser concedida uma vez para cada usuário.",
        "Operações administrativas respeitam o escopo do papel e da organização vinculada ao usuário autenticado.",
        "Credenciais e tokens não são conteúdos de interface nem devem ser inseridos no repositório ou em documentos públicos.",
    ])

    heading(doc, "11. Testes e validação", 1)
    add_para(
        doc,
        "Foram executadas verificacoes de sintaxe nos arquivos JavaScript principais e uma varredura em todos os arquivos JS fora de node_modules. "
        "Tambem foi executada a suite Jest existente, que valida o retorno basico da aplicacao.",
    )
    table([
        ("Verificacao", "Resultado"),
        ("node --check public/script.js", "Aprovado"),
        ("node --check routes/provas.js", "Aprovado"),
        ("node --check config/database.js", "Aprovado"),
        ("node --check public/sw.js", "Aprovado apos correcao"),
        ("Varredura node --check em JS do projeto", "Aprovada"),
        ("npm.cmd run lint", "Aprovado"),
        ("npm.cmd test -- --runInBand", "1 suite / 1 teste aprovado"),
    ], [3400, 5960])

    heading(doc, "11.1 Correções e estabilizações relevantes", 2)
    add_bullets(doc, [
        "A renderização das questões usa escape de HTML e validação antes da montagem das alternativas, evitando falhas de referência e conteúdo inseguro.",
        "O fluxo de simulados persiste a prova ativa antes da exibição e apresenta uma tela de conclusão com desempenho da tentativa.",
        "A prova ao vivo separa salvamento parcial de finalização definitiva, preservando placar e correção coerentes.",
        "A tabela respostas_provas_professor registra a coluna finalizada, e o schema cria participantes e limite de alunos para a prova ao vivo.",
        "O service worker foi revisado para não interceptar áreas autenticadas e evitar cache indevido de respostas administrativas.",
        "Conquistas possuem histórico único por usuário e notificação visual no canto superior direito após novos marcos.",
    ])

    heading(doc, "12. Implantação e operação", 1)
    add_para(
        doc,
        "O projeto pode ser executado localmente com PostgreSQL via Docker Compose ou implantado em ambiente Node.js 20 com PostgreSQL externo. "
        "A configuração é orientada por variáveis de ambiente; segredos não devem ser incluídos em arquivos rastreados, capturas de tela ou documentação compartilhada.",
    )
    add_numbered(doc, [
        "Copiar .env.example para .env e preencher apenas no ambiente local ou no cofre de segredos do provedor.",
        "Subir o PostgreSQL com docker compose up -d postgres ou configurar DATABASE_URL para a instância gerenciada.",
        "Instalar dependências com npm ci em ambientes reproduzíveis ou npm install no desenvolvimento inicial.",
        "Executar migrations com npm run db:migrate antes de liberar uma versão que altere o schema.",
        "Executar npm test e a checagem de sintaxe antes do deploy; iniciar com npm start em produção.",
    ])
    table([
        ("Variavel", "Descricao"),
        ("DATABASE_URL", "String de conexao PostgreSQL."),
        ("DATABASE_SSL", "Ativa SSL quando necessario em producao."),
        ("JWT_SECRET", "Chave de assinatura dos tokens."),
        ("MASTER_PASSWORD", "Senha inicial do usuario master."),
        ("OUTLOOK_EMAIL / OUTLOOK_APP_PASSWORD", "Credenciais do provedor SMTP; usar apenas se SMTP AUTH estiver habilitado para a caixa."),
        ("OPENROUTER_API_KEY / MODEL / URL", "Geracao e correcao por IA."),
    ], [2800, 6560])

    heading(doc, "12.1 Checklist de liberação", 2)
    add_bullets(doc, [
        "Confirmar que DATABASE_URL, JWT_SECRET, APP_URL e variáveis de IA/e-mail estão disponíveis no ambiente de destino.",
        "Executar as migrations e verificar /health/db antes de disponibilizar o acesso público.",
        "Verificar login de aluno, professor e administrador em janela anônima para não confundir sessão com cache local.",
        "Publicar nova versão do service worker quando houver mudança nos ativos e validar que ele não armazena respostas de APIs autenticadas.",
        "Caso o Outlook bloqueie SMTP AUTH, optar por OAuth2/Microsoft Graph ou habilitar o método na caixa corporativa, sem reutilizar senhas pessoais.",
    ])

    heading(doc, "13. Manutenção evolutiva", 1)
    add_para(
        doc,
        "A manutencao deve priorizar migracoes consistentes, testes automatizados para rotas criticas e melhoria gradual da experiencia visual. "
        "A aplicacao ja possui separacao de modulos suficiente para permitir evolucao sem reescrever o sistema inteiro.",
    )
    add_bullets(doc, [
        "Adicionar testes de integracao para /gerar-enem, /salvar-prova, /professor/provas e /provas-prontas.",
        "Criar testes de permissao para todos os perfis administrativos.",
        "Padronizar encoding dos arquivos para UTF-8 e revisar textos com mojibake.",
        "Ampliar observabilidade com logs estruturados por rota e usuario.",
        "Melhorar CI para executar node --check em todos os arquivos JS e rodar Jest com banco de teste.",
        "Adicionar pagina de status para IA/email e tratamentos de fallback.",
    ])

    heading(doc, "14. Conclusão", 1)
    add_para(
        doc,
        "O FormulaVest apresenta uma base funcional para uma plataforma educacional completa, com recursos de estudo individual, prova ao vivo e gestao institucional. "
        "A arquitetura escolhida e objetiva e adequada ao escopo do projeto, enquanto o uso de PostgreSQL, Express e interfaces estaticas facilita implantacao e manutencao.",
    )
    add_para(
        doc,
        "A revisão técnica corrigiu pontos que afetavam diretamente a experiência do aluno e do professor, especialmente a exibição de provas geradas, a conclusão de tentativas, a persistência de conquistas e a navegação autenticada. "
        "Com a documentação, as migrações e os controles operacionais atualizados, o projeto fica mais preparado para evolução, testes e apresentação acadêmica.",
    )

    page_break(doc)
    heading(doc, "Referencias", 1)
    add_bullets(doc, [
        "Documentacao oficial do Node.js e Express para aplicacoes web.",
        "Documentacao do PostgreSQL para modelagem relacional e tipos JSONB.",
        "Documentacao das bibliotecas bcrypt, jsonwebtoken, multer, pdfkit e exceljs.",
        "Arquivos-fonte do projeto FormulaVest analisados nesta revisao.",
    ])

    heading(doc, "Apendice A - Rotas principais", 1)
    table([
        ("Grupo", "Rotas exemplares"),
        ("Autenticacao", "/register, /verificar-email, /login-iniciar, /login-confirmar, /token/refresh, /logout"),
        ("Aluno", "/me, /dashboard, /ranking, /gerar-enem, /gerar-provao, /salvar-prova"),
        ("Prova ao vivo", "/provas-prontas/entrar, /provas-prontas/:id/responder, /provas-prontas/:id/finalizar"),
        ("Professor", "/professor/provas, /professor/provas/:id/iniciar, /professor/provas/:id/placar, /professor/salas"),
        ("Admin", "/admin/usuarios, /admin/criar-aluno, /admin/escolas, /admin/periodos/:id, /admin/salas/:id"),
        ("Master", "/master/stats, /master/empresas, /master/criar-empresa, /master/criar-admin"),
    ], [2300, 7060])

    heading(doc, "Apendice B - Checklist operacional", 1)
    add_bullets(doc, [
        "Executar migrations antes de liberar nova versao.",
        "Limpar cache do service worker quando arquivos publicos mudarem.",
        "Validar login de aluno, professor e admin.",
        "Criar uma prova de professor, transmitir, entrar com aluno, responder e exportar placar.",
        "Gerar ENEM/Provao, responder e confirmar registro no dashboard.",
        "Verificar variaveis de IA e email em producao.",
    ])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_doc())
